"""Tests integración exportación Excel SIP Obra."""

from __future__ import annotations

from io import BytesIO
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIRequestFactory, force_authenticate

from sipo.constants import SIPO_ROL_ADMIN
from sipo.services.excel_export import generar_excel_obras
from sipo.views_excel import SipoExportarExcelView


def _admin_user():
    user = MagicMock()
    user.is_authenticated = True
    user.sip_rol_id = SIPO_ROL_ADMIN
    user.email = 'admin@flesan.cl'
    return user


def _mock_obra():
    obra = MagicMock()
    obra.cf_rrhh_sip_id = 1200
    obra.cf_rrhh_sip_razonsocial = 'Empresa Test'
    obra.cf_rrhh_sip_nombre_uni = 'UN Test'
    obra.cf_rrhh_sip_uni = 'UN01'
    obra.cf_rrhh_sip_cc = 'CC01'
    obra.cf_rrhh_sip_nombre_cc = 'Centro Test'
    obra.cf_rrhh_sip_status = 6
    obra.cf_rrhh_sip_create_user = 'creador@flesan.cl'
    obra.cf_rrhh_sip_adm = 'adm@flesan.cl'
    obra.cf_rrhh_sip_as = 'as@flesan.cl'
    obra.cf_rrhh_sip_create_date = None
    return obra


def _mock_candidato():
    candidato = MagicMock()
    candidato.cf_rrhh_sip_obra_id = 1200
    candidato.cf_rrhh_sip_obra_user_id = '200001'
    candidato.cf_rrhh_sip_obra_candidato_nombre = 'Juan'
    candidato.cf_rrhh_sip_obra_candidato_segundo_nombre = ''
    candidato.cf_rrhh_sip_obra_candidato_ap = 'Pérez'
    candidato.cf_rrhh_sip_obra_candidato_am = 'Gómez'
    candidato.cf_rrhh_sip_obra_candidato_rut = '11111111-1'
    candidato.cf_rrhh_sip_obra_candidato_nomcar = 'Maestro'
    candidato.cf_rrhh_sip_obra_candidato_sueldo = '500000'
    candidato.cf_rrhh_sip_obra_sueldo_base = 450000
    candidato.cf_rrhh_sip_obra_candidato_tipo_contrato = 'Indefinido'
    candidato.cf_rrhh_sip_obra_candidato_seleccionado = 1
    candidato.cf_rrhh_sip_obra_candidato_estado_builder = 1
    candidato.cf_rrhh_sip_obra_candidato_registro_dt = '1'
    return candidato


class SipoExcelExportServiceTests(SimpleTestCase):
    @patch('sipo.services.excel_export.SipoCandidatoObra')
    @patch('sipo.services.excel_export.get_sipo_export_queryset')
    def test_generar_excel_obras_produce_xlsx_valido(self, mock_qs, mock_candidato_model):
        mock_qs.return_value = [_mock_obra()]
        mock_manager = MagicMock()
        mock_manager.only.return_value = mock_manager
        mock_manager.filter.return_value = mock_manager
        mock_manager.order_by.return_value = [_mock_candidato()]
        mock_candidato_model.objects.using.return_value = mock_manager

        buffer = generar_excel_obras({'user': _admin_user(), 'status': 'Activas'})
        content = buffer.getvalue()

        self.assertTrue(content.startswith(b'PK'))
        self.assertGreater(len(content), 0)

        wb = load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ['Solicitudes de Obra', 'Detalle Candidatos'])
        self.assertEqual(wb['Solicitudes de Obra'].max_row, 2)
        self.assertEqual(wb['Detalle Candidatos'].max_row, 2)


class SipoExcelExportViewTests(SimpleTestCase):
    def setUp(self):
        self.factory = APIRequestFactory()

    @patch('sipo.views_excel.generar_excel_obras')
    def test_exportar_excel_ok_headers_y_binario(self, mock_generar):
        from openpyxl import Workbook

        buffer = BytesIO()
        wb = Workbook()
        wb.save(buffer)
        mock_generar.return_value = BytesIO(buffer.getvalue())

        request = self.factory.get('/api/sipo/exportar-excel/', {'status': 'Activas'})
        force_authenticate(request, user=_admin_user())
        response = SipoExportarExcelView.as_view()(request)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment; filename="Reporte_SIPO_Obra.xlsx"', response['Content-Disposition'])
        self.assertTrue(response.content.startswith(b'PK'))
        self.assertGreater(len(response.content), 0)
