from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..models_ficha import SipoFichaIngreso
from .ficha_labels import resolve_catalog_label
from .fichas import apply_ficha_scope, normalize_ficha_estado

FICHA_HEADERS = (
    'ID',
    'Estado',
    'Creado por',
    'Fecha creación',
    'Razón Social ID',
    'Razón Social',
    'Obra',
    'Centro Costo ID',
    'Centro Costo',
    'Cargo',
    'Fecha ingreso',
    'Correo colaborador',
    'Jefe user ID',
    'Jefe nombre',
    'Jefe correo',
    'Correo jefe directo',
    'Correo jefe de terreno',
    'Nombres',
    'Apellido paterno',
    'Apellido materno',
    'Nombre completo',
    'RUT',
    'Tratamiento',
    'Género',
    'AFP',
    'Isapre / Fonasa',
    'Jubilado',
    'Estado civil',
    'Edad',
    'Fecha nacimiento',
    'País nacimiento',
    'Región nacimiento',
    'Nacionalidad',
    'Nacionalidad extranjera',
    'Teléfono',
    'Domicilio',
    'Villa',
    'Número dirección',
    'N° depto',
    'Región',
    'Ciudad',
    'Comuna',
    'Email personal',
    'Banco',
    'Método de pago',
    'N° cuenta',
    'Sueldo líquido',
    'Días contrato',
    'Cuenta de gasto',
    'Tipo contrato',
    'Término contrato / HITO',
    'Fecha término HITO',
    'Tipo jornada',
    'Horario',
    'Observaciones',
    'Doc domicilio',
    'Doc título',
    'Doc AFP',
    'Doc salud',
    'Doc cédula',
    'Aprobado jefe por',
    'Aprobado jefe at',
    'Aprobado RRHH por',
    'Aprobado RRHH at',
    'Rechazo por',
    'Rechazo at',
    'Comentario rechazo',
)

GENERO_LABELS = {
    'M': 'Masculino',
    'F': 'Femenino',
}

ESTADO_ALIASES = {
    'PENDIENTE_RRHH': ('PENDIENTE_RRHH', 'PENDIENTE_ADMIN'),
    'PENDIENTE_JEFE_TERRENO': ('PENDIENTE_JEFE_TERRENO', 'PENDIENTE_JEFE'),
    'PENDIENTE_DATOS_COLABORADOR': (
        'PENDIENTE_DATOS_COLABORADOR',
        'BORRADOR_SUPERVISOR',
    ),
}


def parse_ficha_export_filtros(request) -> dict:
    qp = request.query_params
    return {
        'search': qp.get('search') or qp.get('q') or qp.get('filterText'),
        'estado': qp.get('estado') or qp.get('status'),
        'razon_social_id': qp.get('razon_social_id') or qp.get('empresa'),
        'centro_costo_id': qp.get('centro_costo_id') or qp.get('centro_costo'),
        'user': request.user,
    }


def _format_datetime(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%d-%m-%Y %H:%M')
    if isinstance(value, date):
        return value.strftime('%d-%m-%Y')
    return str(value)


def _genero_label(raw) -> str:
    text = str(raw or '').strip()
    if not text:
        return ''
    return GENERO_LABELS.get(text.upper(), text)


def _bool_label(value) -> str:
    if value is None:
        return ''
    return 'Sí' if bool(value) else 'No'


def _doc_label(field) -> str:
    try:
        return 'Sí' if field and getattr(field, 'name', None) else 'No'
    except (ValueError, FileNotFoundError):
        return 'No'


def _estado_label(ficha: SipoFichaIngreso) -> str:
    code = normalize_ficha_estado(ficha.estado)
    return SipoFichaIngreso.ESTADO_LABELS.get(code, code or '')


def _apply_export_filters(qs, filtros: dict):
    rs = str(filtros.get('razon_social_id') or '').strip()
    if rs:
        rs_clean = rs.replace('.', '').replace(' ', '').upper()
        qs = qs.filter(
            Q(razon_social_id=rs)
            | Q(razon_social_id=rs_clean)
            | Q(razon_social_id__iexact=rs)
            | Q(razon_social_id__iexact=rs_clean)
            | Q(razon_social_nombre__iexact=rs)
        )

    cc = str(filtros.get('centro_costo_id') or '').strip()
    if cc:
        qs = qs.filter(
            Q(centro_costo_id=cc)
            | Q(centro_costo_id__iexact=cc)
            | Q(centro_costo_nombre__iexact=cc)
        )

    estado = str(filtros.get('estado') or '').strip().upper()
    if estado:
        aliases = ESTADO_ALIASES.get(estado, (estado,))
        qs = qs.filter(estado__in=aliases)

    text = str(filtros.get('search') or '').strip()
    if text:
        qs = qs.filter(
            Q(rut__icontains=text)
            | Q(nombres__icontains=text)
            | Q(apellido_paterno__icontains=text)
            | Q(apellido_materno__icontains=text)
            | Q(razon_social_nombre__icontains=text)
            | Q(centro_costo_id__icontains=text)
            | Q(centro_costo_nombre__icontains=text)
            | Q(creado_por__icontains=text)
            | Q(cargo__icontains=text)
            | Q(correo_colaborador__icontains=text)
        )
    return qs


def _ficha_row(ficha: SipoFichaIngreso) -> tuple:
    return (
        ficha.id,
        _estado_label(ficha),
        ficha.creado_por or '',
        _format_datetime(ficha.created_at),
        ficha.razon_social_id or '',
        ficha.razon_social_nombre or '',
        ficha.obra or '',
        ficha.centro_costo_id or '',
        ficha.centro_costo_nombre or '',
        ficha.cargo or '',
        _format_datetime(ficha.fecha_ingreso),
        ficha.correo_colaborador or '',
        ficha.jefe_user_id or '',
        ficha.jefe_nombre or '',
        ficha.jefe_correo or '',
        ficha.correo_jefe_directo or '',
        ficha.correo_admin_obra or '',
        ficha.nombres or '',
        ficha.apellido_paterno or '',
        ficha.apellido_materno or '',
        ficha.nombre_colaborador or '',
        ficha.rut or '',
        ficha.tratamiento or '',
        _genero_label(ficha.genero),
        resolve_catalog_label('afp', ficha.afp) or '',
        resolve_catalog_label('isapre_fonasa', ficha.isapre_fonasa) or '',
        _bool_label(ficha.jubilado),
        resolve_catalog_label('estado_civil', ficha.estado_civil) or '',
        ficha.edad if ficha.edad is not None else '',
        _format_datetime(ficha.fecha_nacimiento),
        ficha.pais_nacimiento or '',
        ficha.region_nacimiento or '',
        ficha.nacionalidad or '',
        ficha.nacionalidad_ext or '',
        ficha.telefono or '',
        ficha.domicilio or '',
        ficha.villa or '',
        ficha.numero_direccion or '',
        ficha.num_depto or '',
        ficha.region or '',
        ficha.ciudad or '',
        ficha.comuna or '',
        ficha.email_personal or '',
        resolve_catalog_label('banco', ficha.banco) or '',
        ficha.metodo_pago or '',
        ficha.numero_cuenta or '',
        ficha.sueldo_liquido if ficha.sueldo_liquido is not None else '',
        ficha.dias_contrato if ficha.dias_contrato is not None else '',
        ficha.cuenta_gasto or '',
        ficha.tipo_contrato or '',
        ficha.termino_contrato or '',
        _format_datetime(ficha.fecha_termino_ito),
        ficha.tipo_jornada or '',
        resolve_catalog_label('horario', ficha.horario) or '',
        ficha.observaciones or '',
        _doc_label(ficha.doc_domicilio),
        _doc_label(ficha.doc_titulo),
        _doc_label(ficha.doc_afp),
        _doc_label(ficha.doc_salud),
        _doc_label(ficha.doc_cedula),
        ficha.aprobado_jefe_por or '',
        _format_datetime(ficha.aprobado_jefe_at),
        ficha.aprobado_rrhh_por or ficha.aprobado_admin_por or '',
        _format_datetime(ficha.aprobado_rrhh_at or ficha.aprobado_admin_at),
        ficha.rechazo_por or '',
        _format_datetime(ficha.rechazo_at),
        ficha.rechazo_comentario or '',
    )


def generar_excel_fichas(filtros: dict) -> BytesIO:
    user = filtros['user']
    qs = apply_ficha_scope(SipoFichaIngreso.objects.all(), user)
    qs = _apply_export_filters(qs, filtros).order_by('-id')
    fichas = list(qs)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Fichas ingreso'
    ws.append(list(FICHA_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for ficha in fichas:
        ws.append(list(_ficha_row(ficha)))

    for idx, _ in enumerate(FICHA_HEADERS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = min(28, max(12, len(FICHA_HEADERS[idx - 1]) + 2))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
