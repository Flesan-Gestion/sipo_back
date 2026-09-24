from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..constants import (
    SIPO_CANDIDATO_ACTIVO,
    SIPO_CANDIDATO_CONTRATADO_SAP,
    SIPO_CANDIDATO_SELECCIONADO,
    SIPO_CANDIDATO_SIN_SELECCION,
    SIPO_ESTADO_LABELS,
)
from ..models import SipoCandidatoObra
from .list import get_sipo_export_queryset

OBRA_HEADERS = (
    'ID Solicitud',
    'Razón Social',
    'Unidad de Negocio',
    'Centro de Costo',
    'Estado',
    'Creador',
    'Administrador de Obra',
    'Asistente',
    'Fecha Creación',
)

CANDIDATO_HEADERS = (
    'ID Solicitud',
    'Correlativo NP',
    'Nombre Completo',
    'RUT',
    'Cargo',
    'Sueldo Líquido',
    'Sueldo Base (Kiptor)',
    'Tipo Contrato',
    'Estado Selección',
    'Estado iBuilder',
    'Revisión DT',
)

CANDIDATO_ONLY_FIELDS = (
    'cf_rrhh_sip_obra_id',
    'cf_rrhh_sip_obra_user_id',
    'cf_rrhh_sip_obra_candidato_nombre',
    'cf_rrhh_sip_obra_candidato_segundo_nombre',
    'cf_rrhh_sip_obra_candidato_ap',
    'cf_rrhh_sip_obra_candidato_am',
    'cf_rrhh_sip_obra_candidato_rut',
    'cf_rrhh_sip_obra_candidato_nomcar',
    'cf_rrhh_sip_obra_candidato_sueldo',
    'cf_rrhh_sip_obra_sueldo_base',
    'cf_rrhh_sip_obra_candidato_tipo_contrato',
    'cf_rrhh_sip_obra_candidato_seleccionado',
    'cf_rrhh_sip_obra_candidato_estado_builder',
    'cf_rrhh_sip_obra_candidato_registro_dt',
)

SELECCION_LABELS = {
    SIPO_CANDIDATO_SIN_SELECCION: 'Sin selección',
    SIPO_CANDIDATO_SELECCIONADO: 'Seleccionado',
    SIPO_CANDIDATO_CONTRATADO_SAP: 'Contratado SAP',
}


def parse_export_filtros(request) -> dict:
    qp = request.query_params
    return {
        'search': qp.get('search') or qp.get('filterText') or qp.get('q'),
        'status': qp.get('status') or qp.get('estado'),
        'centro_costo': qp.get('centro_costo'),
        'empresa': qp.get('empresa') or qp.get('filterEmpresa'),
        'cargo': qp.get('cargo') or qp.get('filterCargo'),
        'fecha_inicio': qp.get('fecha_inicio'),
        'fecha_fin': qp.get('fecha_fin'),
        'user': request.user,
    }


def _format_datetime(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%d-%m-%Y %H:%M')
    if isinstance(value, date):
        return value.strftime('%d-%m-%Y')
    return str(value)


def _estado_label(status) -> str:
    try:
        return SIPO_ESTADO_LABELS.get(int(status), str(status or ''))
    except (TypeError, ValueError):
        return str(status or '')


def _centro_costo_label(obra) -> str:
    code = (obra.cf_rrhh_sip_cc or '').strip()
    name = (obra.cf_rrhh_sip_nombre_cc or '').strip()
    if code and name:
        return f'{code} - {name}'
    return code or name


def _nombre_completo(candidato) -> str:
    parts = (
        candidato.cf_rrhh_sip_obra_candidato_nombre,
        candidato.cf_rrhh_sip_obra_candidato_segundo_nombre,
        candidato.cf_rrhh_sip_obra_candidato_ap,
        candidato.cf_rrhh_sip_obra_candidato_am,
    )
    return ' '.join(str(p).strip() for p in parts if p and str(p).strip())


def _seleccion_label(value) -> str:
    try:
        return SELECCION_LABELS.get(int(value or 0), str(value or ''))
    except (TypeError, ValueError):
        return str(value or '')


def _builder_label(value) -> str:
    return 'Sincronizado' if int(value or 0) == 1 else 'Pendiente'


def _revision_dt_label(value) -> str:
    return 'Sí' if str(value or '').strip() == '1' else 'No'


def _autosize_columns(ws):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)


def _write_header(ws, headers: tuple[str, ...]):
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _obra_row(obra) -> list:
    return [
        obra.cf_rrhh_sip_id,
        obra.cf_rrhh_sip_razonsocial or '',
        obra.cf_rrhh_sip_nombre_uni or obra.cf_rrhh_sip_uni or '',
        _centro_costo_label(obra),
        _estado_label(obra.cf_rrhh_sip_status),
        obra.cf_rrhh_sip_create_user or '',
        obra.cf_rrhh_sip_adm or '',
        obra.cf_rrhh_sip_as or '',
        _format_datetime(obra.cf_rrhh_sip_create_date),
    ]


def _candidato_row(candidato) -> list:
    sueldo_base = candidato.cf_rrhh_sip_obra_sueldo_base
    return [
        candidato.cf_rrhh_sip_obra_id,
        candidato.cf_rrhh_sip_obra_user_id or '',
        _nombre_completo(candidato),
        candidato.cf_rrhh_sip_obra_candidato_rut or '',
        candidato.cf_rrhh_sip_obra_candidato_nomcar or '',
        candidato.cf_rrhh_sip_obra_candidato_sueldo or '',
        str(sueldo_base) if sueldo_base is not None else '',
        candidato.cf_rrhh_sip_obra_candidato_tipo_contrato or '',
        _seleccion_label(candidato.cf_rrhh_sip_obra_candidato_seleccionado),
        _builder_label(candidato.cf_rrhh_sip_obra_candidato_estado_builder),
        _revision_dt_label(candidato.cf_rrhh_sip_obra_candidato_registro_dt),
    ]


def generar_excel_obras(filtros: dict) -> BytesIO:
    user = filtros.get('user')
    obras = list(get_sipo_export_queryset(user, filtros))
    obra_ids = [obra.cf_rrhh_sip_id for obra in obras]

    candidatos = []
    if obra_ids:
        candidatos = list(
            SipoCandidatoObra.objects.using('sip_db')
            .only(*CANDIDATO_ONLY_FIELDS)
            .filter(
                cf_rrhh_sip_obra_id__in=obra_ids,
                cf_rrhh_sip_obra_candidato_estado=SIPO_CANDIDATO_ACTIVO,
            )
            .order_by('cf_rrhh_sip_obra_id', 'cf_rrhh_sip_obra_user_id')
        )

    wb = Workbook()
    ws_obras = wb.active
    ws_obras.title = 'Solicitudes de Obra'
    _write_header(ws_obras, OBRA_HEADERS)
    for obra in obras:
        ws_obras.append(_obra_row(obra))
    _autosize_columns(ws_obras)

    ws_candidatos = wb.create_sheet('Detalle Candidatos')
    _write_header(ws_candidatos, CANDIDATO_HEADERS)
    for candidato in candidatos:
        ws_candidatos.append(_candidato_row(candidato))
    _autosize_columns(ws_candidatos)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
